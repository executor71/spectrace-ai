from pathlib import Path
import sys,json,zipfile
sys.stdout.reconfigure(encoding='utf8')
sys.path.insert(0,str(Path('phase11-browser-tools').resolve()));sys.path.insert(0,str(Path('phase13-tools').resolve()))
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
out=Path('phase13-evidence');result={};base='http://localhost:5173'
with sync_playwright() as p:
 b=p.chromium.launch(channel='chrome',headless=True);context=b.new_context(viewport={'width':1440,'height':1000},accept_downloads=True);page=context.new_page();errors=[];writes=[]
 page.on('pageerror',lambda e:errors.append(str(e)));page.on('request',lambda r:writes.append(r.method+' '+r.url) if r.method not in ['GET','HEAD','OPTIONS'] else None)
 page.goto(base,wait_until='networkidle');page.get_by_role('heading',name='최근 요구사항 목록',exact=True).wait_for()
 response=page.request.get(base+'/api/dashboard/summary');assert response.ok;data=response.json();(out/'final-summary.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
 page.screenshot(path=str(out/'dashboard-desktop.png'),full_page=True)
 assert page.locator('.p13-fill').first.bounding_box()['width']>0
 for width in [1024,390]:
  page.set_viewport_size({'width':width,'height':900 if width==1024 else 844});page.screenshot(path=str(out/f'dashboard-{width}.png'),full_page=True)
  assert not page.evaluate('document.documentElement.scrollWidth>innerWidth'),f'overflow {width}'
 result['responsive']=[1440,1024,390];page.set_viewport_size({'width':1440,'height':1000})
 page.get_by_role('button',name='검토 대기 TC',exact=False).click();assert page.locator('#dashboard-cases select').input_value()=='DRAFT';assert page.locator('.p13-case-list li').count()==data['counts']['DRAFT'];page.screenshot(path=str(out/'filter-draft.png'),full_page=True)
 for key in ['APPROVED','REJECTED','PASS','FAIL','ERROR']:
  label={'APPROVED':'승인','REJECTED':'반려'}.get(key,key);page.get_by_role('button',name=f'{label} {data["counts"][key]}건',exact=True).click();assert page.locator('#dashboard-cases select').input_value()==key
 result['cardAndChartFilters']=True
 target=next(r for r in data['requirements'] if r['counts']['PASS']>0 and r['counts']['total']>=5)
 page.locator('.p13-toolbar select').select_option(target['id']);page.locator('#dashboard-cases select').select_option('ALL');page.screenshot(path=str(out/'filter-requirement.png'),full_page=True)
 report_response=page.request.get(base+f'/api/requirements/{target["id"]}/reports/test-report.xlsx');assert report_response.status==200;assert 'spreadsheetml.sheet' in report_response.headers['content-type'];assert 'filename*=' in report_response.headers['content-disposition']
 with page.expect_download() as download_info:page.locator('#dashboard-requirements').get_by_role('button',name='Excel 다운로드',exact=True).click()
 download=download_info.value;sample=out/download.suggested_filename;download.save_as(str(sample));assert sample.suffix=='.xlsx'
 result['download']={'filename':sample.name,'status':report_response.status,'contentType':report_response.headers['content-type'],'size':sample.stat().st_size,'requirementId':target['id']}
 book=load_workbook(sample,data_only=False);assert book.sheetnames==['요약','요구사항 분석','테스트케이스','실행결과','근거 연결','평가결과']
 summary={row[0].value:row[1].value for row in list(book['요약'].rows)[1:]}
 assert summary['현재 생성 TC']==target['counts']['total'];assert summary['검토 대기']==target['counts']['DRAFT'];assert summary['일반 실행 PASS TC']==target['counts']['PASS']
 for sheet in book:
  assert sheet.freeze_panes=='A2';assert sheet.auto_filter.ref
  for row in sheet:
   for cell in row:assert cell.data_type!='f';assert cell.hyperlink is None
 with zipfile.ZipFile(sample) as z:assert not any('externalLink' in name for name in z.namelist())
 result['xlsxReopen']={'library':'openpyxl 3.1.5','sheets':{s.title:{'rows':s.max_row,'columns':s.max_column} for s in book},'summaryMatchesUI':True,'noFormulasOrExternalLinks':True};book.close()
 page.locator('#dashboard-requirements').get_by_role('button',name='상세 보기',exact=True).click();page.get_by_role('article',name='요구사항 상세').wait_for();page.screenshot(path=str(out/'requirement-detail.png'),full_page=True)
 page.locator('.p13-context-actions').get_by_role('button',name='테스트',exact=True).click();page.get_by_role('article',name='테스트 상세').wait_for();page.screenshot(path=str(out/'testcase-detail.png'),full_page=True)
 page.locator('.p13-context-actions').get_by_role('button',name='실행',exact=True).click();page.get_by_role('heading',name='실행 결과:',exact=False).wait_for();page.screenshot(path=str(out/'execution-detail.png'),full_page=True)
 page.locator('.p13-context-actions').get_by_role('button',name='분析',exact=True).count()
 page.locator('.p13-context-actions').get_by_role('button',name='분석',exact=True).click();page.get_by_role('button',name='저장된 분석 조회',exact=True).click();page.locator('.analysis-result').wait_for();page.screenshot(path=str(out/'analysis-detail.png'),full_page=True)
 # Direct TC navigation selects the specific case rather than the first case.
 page.goto(base,wait_until='networkidle');page.get_by_role('heading',name='최근 요구사항 목록',exact=True).wait_for();page.locator('#dashboard-cases select').select_option('ALL');chosen=data['cases'][-1];page.locator('.p13-case-list').get_by_role('button',name=chosen['scenario'],exact=False).last.click();page.get_by_role('article',name='테스트 상세').wait_for();assert chosen['testCaseId'] in page.get_by_role('article',name='테스트 상세').inner_text();result['specificTCNavigation']=True
 # Historical technical data is loaded from the saved artifact.
 page.goto(base,wait_until='networkidle');page.get_by_role('button',name='기술 검증 결과',exact=True).click();page.get_by_role('cell',name='HYBRID_RERANK',exact=True).wait_for();page.screenshot(path=str(out/'technical-results.png'),full_page=True);result['technicalResults']=True
 # Empty and failure states use intercepted GET responses, never operational DB edits.
 empty={'queriedAt':data['queriedAt'],'requirements':[],'cases':[],'counts':{},'basis':data['basis']}
 page.route('**/api/dashboard/summary',lambda route:route.fulfill(status=200,json=empty));page.goto(base,wait_until='networkidle');page.get_by_text('표시할 요구사항이 없습니다.',exact=True).wait_for();page.screenshot(path=str(out/'empty-state.png'),full_page=True);result['emptyState']=True
 page.unroute('**/api/dashboard/summary');page.route('**/api/dashboard/summary',lambda route:route.fulfill(status=503,json={'message':'검증용 조회 오류'}));page.reload(wait_until='networkidle');page.get_by_role('alert').filter(has_text='검증용 조회 오류').wait_for();page.screenshot(path=str(out/'error-state.png'),full_page=True);result['errorState']=True
 page.unroute('**/api/dashboard/summary');page.reload(wait_until='networkidle');page.get_by_role('heading',name='최근 요구사항 목록',exact=True).wait_for()
 page.route('**/reports/test-report.xlsx',lambda route:route.fulfill(status=500,content_type='application/json',json={'message':'검증용 Excel 오류'}));downloads=[];page.on('download',lambda d:downloads.append(d.suggested_filename));page.locator('#dashboard-requirements').get_by_role('button',name='Excel 다운로드',exact=True).first.click();page.get_by_role('alert').filter(has_text='검증용 Excel 오류').wait_for();assert not downloads;result['errorJsonNotSavedAsXlsx']=True
 result['consoleErrors']=errors;result['mutatingBrowserRequests']=writes;assert not errors;assert not writes
 (out/'browser-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');b.close()
print(json.dumps(result,ensure_ascii=False,indent=2))

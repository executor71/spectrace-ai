from pathlib import Path
import sys,json,io
sys.stdout.reconfigure(encoding='utf8')
sys.path.insert(0,str(Path('phase11-browser-tools').resolve()));sys.path.insert(0,str(Path('phase13-tools').resolve()))
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
out=Path('phase13-evidence');result={}
with sync_playwright() as p:
 b=p.chromium.launch(channel='chrome',headless=True);page=b.new_page(viewport={'width':1440,'height':1000});page.goto('http://localhost:5173',wait_until='networkidle');page.get_by_role('heading',name='최근 요구사항 목록',exact=True).wait_for()
 data=page.request.get('http://localhost:5173/api/dashboard/summary').json();targets=page.request.get('http://localhost:5173/api/evaluations/targets').json();ids={t['requirementId'] for t in targets['requirements']}
 unregistered=next(r for r in data['requirements'] if r['id'] not in ids)
 page.locator('.p13-toolbar select').select_option(unregistered['id']);page.locator('#dashboard-requirements').get_by_role('button',name='상세 보기',exact=True).click();page.locator('.p13-context-actions').get_by_role('button',name='평가',exact=True).click();page.get_by_text('현재 요구사항은 정량평가 대상에 등록되어 있지 않습니다',exact=True).wait_for();page.screenshot(path=str(out/'evaluation-unregistered.png'),full_page=True);result['unregisteredDisplay']=True
 # A registered empty-result state is an injected public GET response, not invented operational evidence.
 registered=targets['requirements'][0];page.route('**/api/requirements/*/evaluations',lambda route:route.fulfill(status=200,json=[]));page.goto('http://localhost:5173',wait_until='networkidle');page.get_by_role('heading',name='최근 요구사항 목록',exact=True).wait_for();page.locator('.p13-toolbar select').select_option(registered['requirementId']);page.locator('#dashboard-requirements').get_by_role('button',name='상세 보기',exact=True).click();page.locator('.p13-context-actions').get_by_role('button',name='평가',exact=True).click();page.get_by_role('button',name='저장된 평가 조회',exact=True).click();page.get_by_text('미평가: 저장된 평가 결과가 없습니다.',exact=False).wait_for();page.screenshot(path=str(out/'evaluation-not-evaluated-mocked.png'),full_page=True);result['notEvaluatedDisplay']={'verified':True,'fixture':'intercepted GET returns [] for registered requirement; DB unchanged'}
 page.unroute('**/api/requirements/*/evaluations')
 response=page.request.get('http://localhost:5173/api/requirements/'+registered['requirementId']+'/reports/test-report.xlsx');assert response.status==200
 (out/'registered-evaluation-report.xlsx').write_bytes(response.body());book=load_workbook(io.BytesIO(response.body()));summary={r[0].value:r[1].value for r in list(book['요약'].rows)[1:]};assert summary['평가 여부']=='저장 평가 있음'
 cells=[str(c.value or '') for sheet in book for row in sheet for c in row]
 assert not any('missingScenarios' in c or 'groundTruthHash' in c or 'HIDDEN_TRUTH' in c for c in cells)
 result['registeredEvaluationReport']={'rows':book['평가결과'].max_row,'scoresPresent':any(isinstance(r[5].value,(float,int)) for r in list(book['평가결과'].rows)[2:]),'hiddenStructuresExcluded':True};book.close()
 (out/'edge-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');b.close()
print(json.dumps(result,ensure_ascii=True))
